from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
R=BASE/'results/full'
A=BASE/'report_assets'
OUT=BASE/'report_build'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.68);s.bottom_margin=Inches(.66)
s.left_margin=Inches(.8);s.right_margin=Inches(.8)
s.header_distance=Inches(.28);s.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.07
normal.paragraph_format.widow_control=True
for name,size in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.space_before=Pt(9 if name!='Title' else 0)
 st.paragraph_format.keep_with_next=True
for name in ['Caption','Subtitle']:
 doc.styles[name].font.size=Pt(9.5 if name=='Caption' else 11)
doc.styles['Caption'].font.italic=True
header=s.header.paragraphs[0];header.text='CMPSC 448    |    LLM Source Attribution';header.runs[0].font.size=Pt(8.5)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Identifying Large Language Models from Their Responses'
doc.core_properties.subject='CMPSC 448 midterm project based on completed full experiments'
doc.core_properties.author=''

def p(text,style=None):
 x=doc.add_paragraph(text,style);return x

def h(text,level=1):return doc.add_heading(text,level)
def page():doc.add_page_break()
def caption(text):return p(text,'Caption')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 pr.append(borders)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();nonsplit=OxmlElement('w:cantSplit');trpr.append(nonsplit)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'65' if side in ['top','bottom'] else '85');el.set(qn('w:type'),'dxa');mar.append(el)
   cp.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'404040' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF'));cp.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1
    if j>0:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs:
     r.font.size=Pt(9.5);r.font.bold=(i==0);r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 p('').paragraph_format.space_after=Pt(0)
 return t

def fig(stem,width=6.25):
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(3);x.paragraph_format.keep_with_next=True
 x.add_run().add_picture(str(A/(stem+'.png')),width=Inches(width))

# Page 1
p('Identifying Large Language Models from Their Responses','Title')
p('CMPSC 448 Midterm Project    |    Individual submission','Subtitle')
p('Author: [Enter your full name before submission]\nReport date: October 10, 2026')
h('Abstract')
p('This project studies whether a small text classifier can identify which of five large language models generated an answer. A balanced subset of LLMRouterBench contains 57,565 responses to 11,513 shared prompts. Prompt groups, including detected near duplicates, remain within one train, validation, or test split. A convolutional neural network (CNN) and a bidirectional LSTM (BiLSTM) are trained from scratch to predict the generator. On the 8,635-response test set, output-only classification reaches mean accuracy of 86.71% for the CNN and 85.67% for the BiLSTM across three seeds. Input-only accuracy is exactly 20%, as required by the paired design; adding the prompt provides no consistent CNN improvement and lowers BiLSTM macro-F1. Holding out a target domain during training reduces CNN macro-F1 from 0.8769 to 0.5187 for coding and from 0.7580 to 0.3599 for factual question answering. Formatting removal lowers performance, while masking selected model and provider names does not yield a detectable reduction. Attribution remains strong on a restricted, fixed-length subset. These results support learnable source cues within the sampled distribution, with substantial dependence on task coverage and response structure.')
h('1 Problem and research questions')
p('The objective is source attribution: given a response, predict its generator label. It is not a comparison of which generator answers questions most correctly. The five source LLMs are not retrained; the trainable systems are two approximately four-million-parameter classifiers. The study addresses the four questions in the course project specification [1].')
for label,text in [
 ('RQ1','Can a CNN and an RNN identify the source model from output text alone?'),
 ('RQ2','How does classification change when the classifier receives only the prompt, only the response, or both?'),
 ('RQ3','Do the learned source cues transfer to a task domain excluded from training?'),
 ('RQ4','Which observable linguistic and structural properties carry source information, and how do formatting, keyword masking, and length controls affect performance?')]:
 x=p('');x.add_run(label+'  ').bold=True;x.add_run(text)
p('The task is closed-set classification among five specified model versions. It does not evaluate unknown generators, human authorship, external editing attacks, or later releases of the same model families.')

# Page 2
page();h('2 Data source and curation')
p('The data are existing model responses released by the authors of LLMRouterBench [2, 3], a benchmark originally developed for model routing. This project repurposes its response records for source classification. All formal experiments use the frozen response release specified in Section 11. No new API responses were generated for this study. The prepared record contains the source label, visible input prompt, original raw output, identifiers, benchmark, domain, and split; only the chosen text fields are fed to the classifier.')
caption('Table 1. Generator labels and exact source identifiers. Each source contributes 11,513 responses.')
table(['Label','Exact model identifier'],[['GPT','gpt-5-chat'],['Claude','claude-sonnet-4'],['Gemini','gemini-2.5-flash'],['Qwen','qwen3-235b-a22b-2507'],['DeepSeek','deepseek-v3-0324']],[1.15,5.75])
p('The starting selection contains 58,340 responses: 11,668 prompts answered by all five sources, from eight benchmarks. Overlapping subset exports are not added again. Tool-use trajectories and grid tasks are excluded. The original raw_output field is used because a processed prediction field can contain only an extracted answer. Correct answers, incorrect answers, short answers, and genuine refusals are retained; answer correctness is not a selection criterion.')
p('Cleaning removes entire five-answer prompt groups to preserve balance: 152 groups contain collection-failure strings from Qwen, and one group has nonidentical visible prompt text across sources. Normalized exact matching then removes two duplicate prompts, leaving 11,513 prompts. Normalization uses Unicode NFKC, case folding, and whitespace collapsing for matching; delivered text is not rewritten. The exclusion of failed Qwen generations conditions the resulting sample on successful collection and may affect representativeness.')
caption('Table 2. Domain composition of the final corpus. Domain balance is not imposed.')
table(['Domain','Prompts','Responses'],[['Knowledge and reasoning','5,223','26,115'],['Factual QA','4,312','21,560'],['Coding','1,303','6,515'],['Mathematics','425','2,125'],['Creative writing','250','1,250'],['Total','11,513','57,565']],[3.7,1.6,1.6])
p('The eight benchmarks are AIME (59 prompts), LiveMathBench (121), LiveCodeBench (1,055), GPQA (198), HLE (2,040), MMLU-Pro (2,985), SimpleQA (4,312), and ArenaHard (743). Domain labels come from benchmark assignments and the source authors’ ArenaHard subsets, not inferred answer style. They are broad task categories rather than mutually exclusive semantic topics.')

# Page 3
page();h('3 Split design and text representation')
p('Splitting individual response rows would allow answers to the same question into both training and testing. Instead, all five answers to a prompt are kept together. Additional connections join original queries of at least 80 characters with character 3-to-5-gram TF-IDF cosine similarity of at least 0.95, and identical normalized responses of at least 50 whitespace words. The resulting 11,461 connected groups are split approximately 70/15/15 with benchmark stratification and data seed 20261005. Lexical grouping reduces known overlap; it cannot certify the absence of every semantic paraphrase.')
caption('Table 3. Fixed splits and class balance.')
table(['Split','Responses','Prompts','Groups','Per source'],[['Train','40,295','8,059','8,007','8,059'],['Validation','8,635','1,727','1,727','1,727'],['Test','8,635','1,727','1,727','1,727']],[1.1,1.5,1.4,1.4,1.5])
p('Audit checks confirm complete five-source prompt groups, identical visible prompts within each group, fixed label mapping, and no shared record, prompt, connected group, or normalized exact input across splits. Recovered copies of all three prepared CSVs match the SHA256 hashes recorded during the formal run. This makes the later integrity review independent of the original saved audit alone.')
p('The tokenizer lowercases Unicode word and digit sequences, retains punctuation and underscores, and represents newlines explicitly. It is a project tokenizer, not an API provider’s tokenizer. The vocabulary contains at most 30,000 tokens including padding, unknown, separator, newline, and identity markers. It is fitted only on the permitted training data. Within a training population, vocabulary construction uses training prompts and response prefixes and is shared across input modes; validation and test text never contribute vocabulary counts.')
caption('Table 4. Input conditions. The response budget is held constant when the prompt is added.')
table(['Condition','Observed tokens'],[['Input only','First 256 prompt tokens'],['Output only','First 512 response tokens'],['Input plus output','First 256 prompt tokens + separator + first 512 response tokens']],[1.6,5.3])
p('Labels, model identifiers, task/domain labels, filenames, and split/group IDs are not concatenated into the classifier text. Original responses remain available for auditing, but neural training sees only the specified prefixes. In training, 34.22% of responses exceed 512 tokens (24.16% to 45.02% across sources), so truncation is a material scope restriction. The maximum combined sequence is 769 tokens. Padding is excluded from CNN pooling windows and BiLSTM recurrent states so batch padding does not become a source cue.')
h('3 1 Why the input only control must score 20 percent',2)
p('Each prompt appears with each of the five source labels. A deterministic classifier receiving only that identical prompt must predict the same class for all five rows. Exactly one prediction is correct, regardless of which class it chooses. Thus 20% accuracy is a structural negative control. Macro-F1 can differ from 0.20 because predictions may be distributed unevenly across classes.')

# Page 4
page();h('4 Models and evaluation protocol')
p('The CNN follows the sentence-classification pattern of convolution and max pooling [4]. The RNN is a bidirectional LSTM [5]. Both learn randomly initialized 128-dimensional embeddings and optimize five-class cross-entropy. No pretrained transformer or external embeddings are used.')
caption('Table 5. Implemented architectures and common training settings.')
table(['Setting','CNN','BiLSTM'],[['Text encoder','1D windows 3, 4, 5; 128 filters each','1 bidirectional layer; 128 hidden units per direction'],['Representation','ReLU, masked max pooling, 384-d concatenation','Concatenated final forward and backward states, 256-d'],['Output','Dropout 0.3; linear five-class layer','Dropout 0.3; linear five-class layer'],['Main parameter count','4,038,917','4,105,477'],['Optimization','Adam, learning rate 0.001; batch 64','Same'],['Stopping','Up to 15 epochs; patience 3 on validation macro-F1','Same'],['Other controls','Gradient norm clip 1.0; deterministic seed settings','Same']],[1.3,2.8,2.8])
p('Adam [6] updates all trainable parameters. The checkpoint with the highest validation macro-F1 is selected, then evaluated on the test set. The configured architecture and training settings are held fixed for the planned comparisons. A maximum-epoch stop is not proof of convergence: the BiLSTM input-plus-output run with seed 42 selects epoch 15, at the training cap.')
p('There are 23 distinct completed results: 14 main neural runs, four cross-domain CNN runs, four additional CNN ablations, and one interpretable-feature logistic regression. Output-only and input-plus-output main conditions use seeds 42, 43, and 44. The two input-only controls, all cross-domain experiments, and all ablations use seed 42. The main CNN output-only seed-42 fit is reused as the full-test ablation reference rather than counted as another run.')
h('4 1 Metrics and uncertainty',2)
p('Accuracy is the proportion of correctly attributed responses. For each source, precision and recall determine its F1; macro-F1 is the unweighted mean of the five source-specific F1 scores, with zero used for undefined cases. Class balance makes accuracy easy to interpret, but macro-F1 also exposes uneven class failures. Per-class and per-domain metrics accompany the overall scores.')
p('For the four main conditions containing response text, results report the mean and sample standard deviation across three training seeds. Confidence intervals use 500 percentile bootstrap resamples of whole test group IDs, retaining all five paired responses. Paired deltas use identical test records in each arm and resample the same groups. These intervals condition on fitted models; they do not include training-seed uncertainty. The several intervals are not adjusted for multiple comparisons. Additional one-seed experiments are sensitivity analyses, not evidence of universal effects.')
p('An independent review reconstructed all 23 accuracy and macro-F1 values from saved per-example predictions, checked record alignment and data hashes, and inspected the run configurations and manifests. This is an internal-consistency audit of saved evidence, without reloading weights or independently rerunning training or inference. Complete per-run scores and stopping epochs appear in Appendix A.')

# Page 5
page();h('5 RQ1 and main classification results')
main=pd.read_csv(R/'tables/main_results.csv')
rows=[]
for arch,lab in [('cnn','CNN'),('bilstm','BiLSTM')]:
 for mode,title in [('input_only','Input only'),('output_only','Output only'),('input_and_output','Input plus output')]:
  z=main[(main.architecture==arch)&(main['mode']==mode)].iloc[0]
  fmt=lambda k,d:(f'{z[k+"_mean"]*100:.2f} ± {z[k+"_std"]*100:.2f}' if k=='accuracy' else f'{z[k+"_mean"]:.4f} ± {z[k+"_std"]:.4f}') if z[k+'_count']>1 else (f'{z[k+"_mean"]*100:.2f}' if k=='accuracy' else f'{z[k+"_mean"]:.4f}')
  rows.append([lab,title,str(int(z.accuracy_count)),fmt('accuracy',2),fmt('macro_f1',4)])
caption('Table 6. Test results on the same 8,635 responses. ± denotes across-seed SD, not a confidence interval.')
table(['Model','Input condition','Seeds','Accuracy %','Macro F1'],rows,[.8,1.75,.55,1.7,2.1])
fig('fig01_main_macro_f1',6.2)
caption('Figure 1. Mean macro-F1 and individual seed results. Input-only has one seed; each response-bearing condition has three. Macro-F1 is not constrained to the 20% input-only accuracy.')
p('Output text alone provides a strong source signal: CNN accuracy is 86.71% and BiLSTM accuracy is 85.67%, compared with the balanced-class accuracy baseline of 20%. CNN output-only macro-F1 is higher by 0.0097 on average, with the same direction in all three seeds. This is a modest advantage under the implemented architectures and training budget, not a general claim that CNNs outperform recurrent models.')
p('Both input-only controls attain exactly 20% accuracy, consistent with the paired design. Their lower and different macro-F1 values reflect prediction imbalance rather than a violation of the control. Neither the input-only scores nor the attribution scores evaluate whether the original LLM answers are factually correct.')

# Page 6
page();h('5 1 Errors differ by source and task',2)
fig('fig04_main_confusion',4.75)
caption('Figure 2. Row-normalized CNN output-only confusion matrix, seed 42. Each source has 1,727 test responses. Diagonal values are recalls. Seed 42 is shown for a fixed illustrative fit, not selected as the best seed.')
p('Across three CNN seeds, class F1 is highest for Claude (0.9330) and DeepSeek (0.9133), followed by Gemini (0.8506), GPT (0.8425), and Qwen (0.7942). In Figure 2, 166 GPT and 148 Gemini responses are assigned to Qwen. These asymmetric confusions indicate overlapping response patterns; they do not establish a particular causal style feature or a source-model quality ranking.')
dom=pd.read_csv(R/'tables/domain_results.csv')
domain_order=[('knowledge_reasoning','Knowledge and reasoning',3920),('factual_qa','Factual QA',3235),('coding','Coding',990),('math','Mathematics',300),('creative_writing','Creative writing',190)]
rows=[]
for key,name,n in domain_order:
 vals=[]
 for arch in ['cnn','bilstm']:
  z=dom[(dom.rq=='main')&(dom.architecture==arch)&(dom['mode']=='output_only')&(dom.task_domain==key)].iloc[0]
  vals.append(f'{z.macro_f1_mean:.4f} ± {z.macro_f1_std:.4f}')
 rows.append([name,f'{n:,}',*vals])
caption('Table 7. Output-only domain macro-F1, mean ± SD over three seeds.')
table(['Domain','Test rows','CNN','BiLSTM'],rows,[2.1,.8,2,2])
p('The aggregate result is dominated by knowledge/reasoning and factual QA; it is not an equally weighted task average. CNN is strongest overall but BiLSTM is better on coding and creative writing. Creative writing is the most difficult domain and contains only 38 test prompts, so its estimates are especially uncertain. The gap between knowledge/reasoning and factual QA motivates the transfer experiment rather than treating a single overall number as sufficient.')

# Page 7
page();h('6 RQ2 Effect of adding the prompt')
p('Adding the prompt preserves the same first 512 response tokens. Therefore, any measured difference is not caused by allocating less response text to make room for the prompt. The paired comparison is input-plus-output minus output-only, using the same architecture, seed, and test records.')
rq2=pd.read_csv(R/'tables/rq2_paired_deltas.csv')
caption('Table 8. Paired changes in macro-F1. Intervals are conditional 95% group-bootstrap intervals.')
rows=[]
for _,z in rq2.iterrows():rows.append(['CNN' if z.architecture=='cnn' else 'BiLSTM',int(z.seed),f'{z.macro_f1_delta:+.4f}',f'[{z.macro_f1_delta_ci_low:+.4f}, {z.macro_f1_delta_ci_high:+.4f}]'])
table(['Model','Seed','Change','95% interval'],rows,[1.25,.65,1.7,3.3])
p('For CNN, mean macro-F1 changes from 0.8667 to 0.8648, a difference of -0.0019. The sign varies across seeds, and two of the three intervals include zero. The appropriate conclusion is that this experiment detects no robust benefit from concatenating the prompt. It does not prove that a different architecture or context representation could never benefit.')
p('For BiLSTM, macro-F1 falls from 0.8571 to 0.8477, a mean change of -0.0093. All three fitted-seed intervals are below zero. The longer combined sequence increases the burden on the final recurrent representation and on optimization, which are plausible explanations rather than tested mechanisms. The seed-42 combined-input run reaches the 15-epoch cap, so the conclusion is conditional on this fixed training budget.')
h('6 1 Computational cost',2)
caption('Table 9. Mean recorded runtime per main response-bearing run in the recorded CPU environment.')
table(['Architecture','Output only','Input plus output'],[['CNN','22.23 minutes','35.56 minutes'],['BiLSTM','458.15 minutes','1,349.44 minutes']],[2.1,2.4,2.4])
p('The recorded environment is Windows 11, Python 3.12.10, PyTorch 2.14.1+cpu, four CPU threads, and no available CUDA device. The 22 neural runs total approximately 96.68 recorded hours. The timing field includes training and the initial test pass, but excludes later bootstrap/report generation; it is not a complete end-to-end stopwatch or a controlled hardware benchmark.')
p('The two main architectures have similar parameter counts, yet recurrent processing is substantially slower in this implementation. Combined with its slightly stronger overall accuracy, the output-only CNN offers the better observed accuracy/runtime tradeoff. This practical conclusion is specific to the tested code, machine, and stopping behavior.')

p('Saved histories also justify validation-based selection. For CNN output-only seed 42, training cross-entropy falls from 0.2225 at selected epoch 7 to 0.1510 at epoch 10, while validation loss rises from 0.3722 to 0.4249. The reported test result uses epoch 7, preserving the validation-selected checkpoint despite continued training-loss improvement.')

# Page 8
page();h('7 RQ3 Transfer to unseen task domains')
p('For each target domain, two output-only CNNs are trained with seed 42. The held-out condition removes every connected training and validation group that contains the target domain. The matched-mixed condition may include target-domain training examples but is downsampled to the same number of training prompts. Both conditions use the same non-target validation set and identical target-domain test set. Their vocabularies are fitted separately from permitted training text.')
caption('Table 10. Cross-domain design counts. Each prompt contributes five responses.')
table(['Target','Train prompts per arm','Target prompts in mixed train','Validation prompts','Test prompts'],[['Coding','7,140','797','1,540','198'],['Factual QA','5,041','1,887','1,080','647']],[1.0,1.55,1.7,1.35,1.3])
fig('fig02_cross_domain',6.25)
caption('Figure 3. Matched-mixed versus target-held-out macro-F1. Error bars are conditional 95% group-bootstrap intervals; paired-difference intervals are reported below.')
p('Coding macro-F1 falls from 0.8769 to 0.5187: change -0.3582, 95% interval [-0.3955, -0.3245]. Factual-QA macro-F1 falls from 0.7580 to 0.3599: change -0.3981, interval [-0.4179, -0.3792]. Accuracy falls from 87.58% to 57.68% for coding and from 75.64% to 37.71% for factual QA. The decline cannot be attributed simply to a smaller number of training prompts because that count is matched.')
p('Errors are uneven across sources. On held-out coding, Qwen recall is 12.12% and GPT recall is 16.16%; 151 of 198 Qwen answers are predicted as Claude. On held-out factual QA, Qwen recall is 11.75% and DeepSeek recall is 17.77%. Some source information transfers, but the large reductions show that within-distribution accuracy does not imply a domain-invariant generator signature.')
p('This experiment covers two broad target domains, one CNN architecture, and one seed. Its intervals do not measure variability from training-sample selection or other seeds. It changes task and benchmark exposure together; it cannot separate every aspect of domain shift from collection-specific response patterns.')

# Page 9
page();h('8 RQ4 Formatting keywords and interpretable features')
p('The full-test ablations retrain output-only CNNs with seed 42 after transforming the original first-512-token span. Deleted tokens are not replaced with later content. These comparisons measure learnability under each representation, rather than edited-input sensitivity of one fixed classifier.')
caption('Table 11. Full-test ablations on 8,635 identical test records. Reference macro-F1 is 0.8652.')
table(['Condition','Macro F1','Change from raw','95% interval'],[['Original prefix','0.8652','Reference','-'],['Formatting removed','0.8283','-0.0369','[-0.0434, -0.0297]'],['Model keywords masked','0.8686','+0.0034','[-0.0019, +0.0087]']],[2.15,1.0,1.5,2.25])
p('Removing approximate Markdown structures, code-fence delimiters, list/heading markers, and whitespace layout lowers macro-F1 by 0.0369 (3.69 points on a 0-to-100 scale). Structure is informative, while substantial signal remains. The rules also change token counts and may remove syntax or content, preventing a pure causal interpretation of formatting.')
p('Masking replaces the five model names and covered provider variants with one identity marker, independently of the true label. No detectable loss appears; indirect identity clues may remain. Matches occur in only 168 of 40,295 training prefixes (0.417%), limiting expected aggregate impact, and need not be self-identifications.')
h('8 1 A transparent feature baseline',2)
p('Multinomial logistic regression uses 13 response features: word and character counts, mean word length, approximate sentence length, paragraph count, heading/list rates, code-fence count, question/exclamation rates, digit and punctuation rates, and lexical type-token ratio (TTR). All predictive features use the same original response prefix available to the neural baseline. Median imputation, a missing-value indicator, and scaling are fitted on training data only. C is fixed at 1.0 rather than tuned on the test set.')
p('The feature classifier achieves 52.58% accuracy and 0.5247 macro-F1 (validation macro-F1 0.5141), showing that simple observable properties carry source information. Neural models outperform this linear feature baseline, but the gain could reflect richer text information, nonlinear combinations of cues, or both. Correlated feature coefficients are associations, not independent importance scores or explanations of the CNN’s internal decisions.')
caption('Table 12. Training-only descriptions. Full-response lengths and paired first-128-word TTR have different scopes.')
feat=pd.read_csv(R/'descriptive/tables/training_features_by_model.csv').set_index('LLM_name')
ttr=pd.read_csv(R/'descriptive/tables/training_paired_ttr128_summary.csv').set_index('LLM_name')
rows=[[n,f'{feat.loc[n,"word_count__mean"]:.1f}',f'{feat.loc[n,"word_count__median"]:.0f}',f'{ttr.loc[n,"mean"]:.3f}'] for n in ['GPT','Claude','Gemini','Qwen','DeepSeek']]
table(['Source','Mean words','Median words','Paired TTR'],rows,[1.4,1.6,1.6,2.3])
p('Length is strongly right-skewed, especially for Gemini and Qwen. TTR uses the first 128 whitespace words on 3,622 training prompts for which all five sources are eligible; tokens that become empty after normalization are dropped. The resulting mean denominators vary from 108.73 for GPT to about 115-119 for the others; higher TTR therefore does not establish richer vocabulary. Full-response lengths are descriptive, not the prefix-limited classifier inputs.')

# Page 10
page();h('8 2 Controlling observed sequence length')
p('The fixed-length experiment chooses the largest K in {32, 64, 128, 256} that retains at least 30% and at least 100 training prompts, using training lengths only. K = 256 is selected. A connected group is eligible only when every response has at least K tokenizer tokens. The same rule is then applied to validation and test. Two fresh CNNs compare the original prefix of up to 512 tokens with exactly the first 256 tokens on identical retained records.')
caption('Table 13. Retention under the fixed-length eligibility rule.')
table(['Split','Retained prompts','Retained responses','Prompt retention'],[['Train','3,281','16,405','40.71%'],['Validation','660','3,300','38.22%'],['Test','704','3,520','40.76%']],[1.2,1.8,1.9,2.0])
fig('fig03_ablation_panels',6.25)
caption('Figure 4. The two panels represent different training and test populations. Only within-panel contrasts are controlled comparisons. Error bars are conditional 95% group-bootstrap intervals.')
p('On the long-answer subset, macro-F1 changes from 0.9605 for the original prefix to 0.9514 for exactly 256 tokens. The paired change is -0.0091, with interval [-0.0160, -0.0024]. Strong classification remains when observed tokenizer sequence length is constant, supporting source information beyond variation in that length within this selected population.')
p('Selection materially changes the task mixture: 2,560 of 3,520 retained test responses are knowledge/reasoning (72.73%), while only five are factual QA (one prompt), compared with 3,235 factual-QA responses in the full test set. Consequently, the approximately 95% subset score cannot be interpreted as an improvement over the approximately 87% full-test score. Nor is the five-row factual-QA subset large enough to support a domain-specific conclusion.')
p('Truncation removes content as well as length variation. Equal tokenizer-token counts do not equalize word counts, characters, formatting, semantics, or full original response length. This experiment is a conditional sensitivity check, not a causal estimate of how much of the full-data performance is explained by verbosity.')

# Page 11
page();h('9 Concrete errors and limits of interpretation')
p('Post hoc error inspection joins saved predictions to the verified test CSV by record_id. Examples below illustrate failure modes after evaluation; they are not used to filter the test set or choose a new model. The accompanying case file preserves full identifiers for traceability.')
h('9 1 Identical answers can make attribution ambiguous',2)
p('For a SimpleQA question about the founding year of Cáceres in Antioquia, Gemini and Qwen produce exactly the same response: “The municipality of Cáceres, Antioquia, Colombia, was founded in **1576**.” The seed-42 CNN predicts Gemini for both, with the same maximum softmax value of 0.7852. Thus Qwen’s answer is misclassified even though no text-only deterministic classifier can distinguish these two identical strings. Shared short factual answers are a real limitation of attribution, not necessarily a training defect.')
p('In another SimpleQA case concerning the cemetery where Camille Pissarro’s youngest son was buried, a Qwen response is classified as GPT with maximum softmax 0.9949. A large softmax value does not ensure a correct source attribution. These probabilities were not calibrated, and the project does not validate a confidence threshold for real-world decisions.')
h('9 2 Scope and threats to validity',2)
p('The five labels represent one specified model version/configuration each. Provider settings, prompt templates, refusal behavior, answer formatting, and collection choices may all influence classification. Visible prompt equality does not establish identical hidden system instructions or decoding settings. The model can learn collection-specific regularities rather than an immutable family fingerprint.')
p('Prompt-level splitting and lexical near-duplicate grouping address known leakage routes, but do not guarantee complete semantic independence. The benchmark mixture is uneven, and a response-level aggregate can hide weak performance on smaller or shifted domains. No later collection period, unseen generator, human-text class, or translation condition is tested. There is no robustness evaluation of a fixed classifier on externally edited or paraphrased outputs.')
p('Three seeds characterize only the main response-bearing comparisons. Cross-domain and ablation experiments use one seed and conditional test-group intervals; these do not capture all retraining variability. Approximate linguistic features behave differently for prose, mathematics, and code. Results from short or highly selected subsets require especially narrow interpretation.')
h('10 Lessons learned and next steps')
p('The central methodological lesson is that a useful classifier score depends on the evaluation design. Pairing the same question across generators makes class balance meaningful and creates an exact negative control. Grouped splits reduce question leakage. Holding the output budget and test records fixed makes the input and ablation comparisons interpretable. The large cross-domain losses show why random held-out test performance alone is insufficient.')
p('The practical lesson is to plan computation before scaling experiments. Similar parameter counts do not imply similar runtime: recurrent sequence processing dominates the recorded CPU budget. A short smoke run, epoch timing, validation-based stopping, saved histories, and restartable checkpoints make long runs manageable. The resulting logs support this report without another training cycle.')
p('A follow-up study should first replicate domain and ablation effects across more seeds, then evaluate a fresh collection with matched generation settings and additional versions per family. Longer-context alternatives, probability calibration, and unknown-source rejection would address limitations of the current closed-set classifier. Those experiments are future work, not completed results.')

# Page 12
page();h('11 Reproducibility and conclusion')
p('The repository package contains the Python data-preparation and training code, fixed configuration, tests, environment records, per-run metrics and predictions, aggregate tables, selected figures, and an independent verification script. The data archive is externally downloadable; processed text and trained checkpoint files are not bundled in the compact repository package. The full prepared data can be recreated from the pinned source revision and checked against the saved hashes. DATA.md and README.md give the exact commands and output locations.')
p('The frozen dataset revision is 0e5af1b84bf73437a01a1849c0f1d2468baa93fc. The source archive SHA256 is b79f8cde1a6f029c2efa663a3a3b6f7748defb22341fe59f328cebef6648c8f1. Run configurations retain architecture, seed, text budgets, transformations, and checkpoint-selection information. Saved predictions permit recalculation of every reported accuracy and macro-F1 without rerunning model training. Fresh inference requires a retained local checkpoint or a reproducible retraining run.')
p('The results answer all four research questions: output-only attribution is strong within this benchmark mixture; prompts alone carry no label information under the paired design and concatenating them does not improve the tested models overall; cross-domain transfer is substantially weaker; and observable response features support above-chance attribution, with sensitivity to formatting normalization. The most defensible conclusion is successful attribution within a defined data collection, with limited evidence for general-purpose or domain-invariant model identification.')
h('11 1 Assistance and author verification',2)
p('ChatGPT/Codex assisted with study design, Python implementation and debugging, result checks, interpretation, and drafting this report. The formal experiments were run by the student in a local Windows environment. This disclosure describes the assistance used; it does not replace the student’s responsibility to verify the analysis, understand the code, and explain the submitted work under the course policy [1].')
h('References')
refs=[
 '[1] CMPSC 448. Midterm Project: Who Wrote It? Identifying LLMs from Their Responses. Course project specification, modified September 17, 2026. Instructor-provided PDF.',
 '[2] H. Li et al. LLMRouterBench: A Massive Benchmark and Unified Framework for LLM Routing. arXiv:2601.07206, 2026. https://arxiv.org/abs/2601.07206',
 '[3] LLMRouterBench authors. Code and response dataset. https://github.com/ynulihao/LLMRouterBench ; https://huggingface.co/datasets/NPULH/LLMRouterBench . Frozen data revision stated above. Accessed October 10, 2026.',
 '[4] Y. Kim. Convolutional Neural Networks for Sentence Classification. Proceedings of EMNLP, pp. 1746-1751, 2014. https://doi.org/10.3115/v1/D14-1181',
 '[5] S. Hochreiter and J. Schmidhuber. Long Short-Term Memory. Neural Computation, 9(8), pp. 1735-1780, 1997. https://doi.org/10.1162/neco.1997.9.8.1735',
 '[6] D. P. Kingma and J. Ba. Adam: A Method for Stochastic Optimization. ICLR, 2015. https://arxiv.org/abs/1412.6980'
]
for ref in refs:
 x=p(ref);x.paragraph_format.space_after=Pt(5)
 for run in x.runs:run.font.size=Pt(10)

# Page 13
page();h('Appendix A Complete experiment inventory')
p('Every row is a distinct completed result. Accuracy is shown as a percentage; macro-F1 is on a 0-to-1 scale. Best/total gives the selected validation epoch and completed training epochs. Minutes use each run’s recorded timing field. Single-run results are not averaged across seeds.')
allruns=pd.read_csv(R/'tables/all_run_metrics.csv')
rows=[]
for _,z in allruns.iterrows():
 if z.rq=='main':
  nam=('CNN' if z.architecture=='cnn' else 'BiLSTM')+' '+{'input_only':'input','output_only':'output','input_and_output':'both'}[z['mode']]
 elif z.rq=='rq3':nam=('Coding' if z.target=='coding' else 'Factual QA')+' '+('mixed' if z.condition=='matched_mixed' else 'held out')
 else:nam={'no_format':'CNN no format','keyword_mask':'CNN keyword mask','length_subset_raw':'CNN long subset raw','fixed_token_length':'CNN fixed 256','fixed_interpretable_features':'Feature logistic'}[z.condition]
 epochs='-' if pd.isna(z.best_epoch) else f'{int(z.best_epoch)}/{int(z.epochs_completed)}'
 secs='-' if pd.isna(z.seconds) else f'{z.seconds/60:.1f}'
 rows.append([nam,int(z.seed),f'{int(z.n_test):,}',f'{z.accuracy*100:.2f}',f'{z.macro_f1:.4f}',epochs,secs])
table(['Experiment','Seed','Test n','Acc %','F1','Best/total','Minutes'],rows,[2.25,.45,.7,.7,.7,.7,1.4])
p('Main runs use 40,295 training responses and 8,635 validation responses. Cross-domain training and validation sizes follow Table 10 (multiply prompts by five). Both length-subset runs use 16,405 training and 3,300 validation responses. Other full-test ablations and the feature baseline use the original splits. The original CNN output seed-42 row also serves as the reference for no-format and keyword masking; it is not an additional fit.')
p('Evidence map: tables/main_results.csv and all_run_metrics.csv support the main and inventory tables; rq2_paired_deltas.csv, rq3_paired_deltas.csv, and rq4_paired_deltas.csv support paired intervals; rq3_design and rq4_design specify matched populations; per-run predictions and confusion matrices support error analysis; descriptive/tables contains training-only linguistic summaries. The independent audit and case-provenance files accompany the package.')

# Remove inherited theme font overrides and decorative paragraph borders.
for style in doc.styles:
 for el in list(style.element.iter(qn('w:pBdr'))):
  el.getparent().remove(el)
 if style.type == 1 and style.name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
  style.font.name='Times New Roman'
  fonts=style.element.find('.//' + qn('w:rFonts'))
  if fonts is not None:
   for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
    fonts.attrib.pop(qn('w:'+key),None)
for par in doc.paragraphs:
 for el in list(par._p.iter(qn('w:pBdr'))):
  el.getparent().remove(el)
path=OUT/'Midterm_Report_Draft.docx'
doc.save(path)
print(path)
